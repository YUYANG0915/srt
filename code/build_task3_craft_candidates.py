#!/usr/bin/env python3
"""Build Task3 P6 craft / art / creation-order candidate review pack from COCO val2017."""

from __future__ import annotations

import argparse
import csv
import html
import io
import json
import ssl
import time
import urllib.request
import zipfile
from collections import defaultdict
from dataclasses import dataclass
from pathlib import Path

try:
    import certifi
except ModuleNotFoundError:
    certifi = None


ROOT = Path(__file__).resolve().parents[1]
DATA_DIR = ROOT / "data"
ZIP_PATH = DATA_DIR / "coco_annotations" / "annotations_trainval2017.zip"
IMAGE_DIR = DATA_DIR / "task3_craft_coco_images"
SSL_CONTEXT = ssl.create_default_context(cafile=certifi.where()) if certifi else ssl.create_default_context()


@dataclass
class Candidate:
    image_id: int
    image_path: str
    captions: list[str]
    categories: list[str]
    suggested_label: str
    suggested_boundary_family: str
    suggested_actor: str
    suggested_tool: str
    suggested_target: str
    suggested_state_cues: str
    suggestion_reason: str
    trigger_words: list[str]
    priority_score: int
    subfamily: str


@dataclass
class CandidateMeta:
    image_id: int
    file_name: str
    captions: list[str]
    categories: list[str]
    suggested_label: str
    suggested_boundary_family: str
    suggested_actor: str
    suggested_tool: str
    suggested_target: str
    suggested_state_cues: str
    suggestion_reason: str
    trigger_words: list[str]
    priority_score: int
    subfamily: str


def load_coco_metadata() -> tuple[dict[int, list[str]], dict[int, list[str]], dict[int, str]]:
    with zipfile.ZipFile(ZIP_PATH) as zf:
        caps = json.load(io.TextIOWrapper(zf.open("annotations/captions_val2017.json"), encoding="utf-8"))
        inst = json.load(io.TextIOWrapper(zf.open("annotations/instances_val2017.json"), encoding="utf-8"))

    cat_name = {row["id"]: row["name"] for row in inst["categories"]}
    caps_by_img: dict[int, list[str]] = defaultdict(list)
    cats_by_img: dict[int, list[str]] = defaultdict(list)
    file_by_img: dict[int, str] = {}
    for image in inst["images"]:
        file_by_img[image["id"]] = image["file_name"]
    for ann in caps["annotations"]:
        caps_by_img[ann["image_id"]].append(ann["caption"])
    for ann in inst["annotations"]:
        cats_by_img[ann["image_id"]].append(cat_name[ann["category_id"]])
    return caps_by_img, cats_by_img, file_by_img


def contains_any(text: str, terms: list[str]) -> list[str]:
    return [term for term in terms if term in text]


def has_any(text: str, terms: list[str]) -> bool:
    return any(term in text for term in terms)


def infer_subfamily(text: str, categories: list[str]) -> str:
    cats = set(categories)
    if has_any(text, ["frosting", "icing", "decorate", "decorating"]):
        return "cake_decoration"
    if "cake" in cats or "cake" in text:
        return "cake_cutting_or_finished"
    if has_any(text, ["drawing", "draw", "painting", "paint", "coloring", "art", "artist"]):
        return "drawing_painting"
    if has_any(text, ["writing", "write"]) or "book" in cats:
        return "writing_paper"
    if has_any(text, ["scissors", "cutting", "cut ", "paper"]) or "scissors" in cats:
        return "cutting_paper"
    if has_any(text, ["making", "assembling", "constructing", "sewing"]):
        return "making_assembly"
    return "other_creation_salience"


def infer_candidate(captions: list[str], categories: list[str]) -> tuple[str, str, str, str, str, str, str, list[str], int] | None:
    text = " ".join(captions).lower()
    cats = set(categories)

    triggers = contains_any(
        text,
        [
            "drawing",
            "draw",
            "painting",
            "paint",
            "writing",
            "write",
            "coloring",
            "craft",
            "making",
            "decorate",
            "decorating",
            "cutting",
            "scissors",
            "paper",
            "art",
            "artist",
            "cake",
            "frosting",
            "icing",
            "sewing",
            "assembling",
            "constructing",
        ],
    )
    for cat in ["scissors", "cake", "book", "teddy bear"]:
        if cat in cats:
            triggers.append(cat)
    if not triggers:
        return None

    score = len(set(triggers))
    if "person" in cats:
        score += 3
    if has_any(text, ["drawing", "painting", "writing", "cutting", "frosting", "icing", "decorating", "sewing", "assembling"]):
        score += 7
    if {"scissors", "cake"} & cats:
        score += 3

    finished_art_phrases = ["painting of", "a painting", "an image of", "art work", "artwork", "painted picture", "drawing of"]
    if has_any(text, finished_art_phrases):
        return (
            "finished_artifact",
            "active_vs_finished",
            "none_or_unclear",
            "none_or_unclear",
            "painting/drawing/art object",
            "finished artwork is visible without the act of making it",
            "caption describes an artwork as an object, not an active creation process",
            triggers,
            score + 2,
        )

    if has_any(text, ["drawing", "draw", "painting", "paint", "coloring", "writing", "write"]):
        if has_any(text, ["drawing", "coloring", "writing"]) or has_any(text, ["painting on", "painting a", "painting with"]):
            return (
                "active_creation",
                "setup_vs_active",
                "person/hand",
                "pen/brush/marker/pencil",
                "paper/canvas/book/surface",
                "hand/tool appears engaged with a creation surface",
                "caption suggests active mark-making or writing",
                triggers,
                score + 7,
            )
        return (
            "pre_creation_setup",
            "setup_vs_active",
            "person/hand",
            "pen/brush/marker/pencil",
            "paper/canvas/book/surface",
            "art/writing materials may be present but contact needs visual checking",
            "caption has weak art/writing cue but may be setup or result",
            triggers,
            score + 3,
        )

    if has_any(text, ["cutting", "cut ", "scissors"]):
        if has_any(text, ["cutting", "cut "]):
            return (
                "active_creation",
                "setup_vs_active",
                "person/hand",
                "scissors/knife",
                "paper/food/material",
                "tool-target contact suggests active cutting stage",
                "cutting is an active construction or preparation step",
                triggers,
                score + 6,
            )
        return (
            "pre_creation_setup",
            "setup_vs_active",
            "person/hand",
            "scissors",
            "paper/material",
            "scissors are visible but active cutting must be checked",
            "scissors presence may be setup rather than active creation",
            triggers,
            score + 3,
        )

    if has_any(text, ["being decorated", "decorating", "frosting a cake", "icing a cake"]):
        return (
            "active_creation",
            "assembly_vs_finished",
            "person/hand",
            "icing/frosting tool/hand",
            "cake/dessert",
            "decoration material is being applied or manipulated",
            "decorating/icing is an active finishing stage before final presentation",
            triggers,
            score + 6,
        )

    if "cake" in cats or "cake" in text:
        if has_any(text, ["cutting", "cuts", "cut into", "slicing", "about to cut"]):
            return (
                "active_creation",
                "setup_vs_active",
                "person/hand",
                "knife/cake cutter",
                "cake/dessert",
                "knife is positioned at or inside the cake",
                "cake cutting is an active serving/modification step",
                triggers,
                score + 5,
            )
        return (
            "finished_artifact",
            "active_vs_finished",
            "none_or_unclear",
            "none_or_unclear",
            "cake/dessert",
            "finished cake is visible without clear active creation contact",
            "cake scenes often show final artifact rather than making process",
            triggers,
            score + 2,
        )

    if has_any(text, ["making", "assembling", "constructing", "sewing"]):
        return (
            "active_assembly",
            "setup_vs_active",
            "person/hand",
            "hands/tools",
            "object/material",
            "caption suggests object construction or assembly in progress",
            "making/building cue indicates an active construction process",
            triggers,
            score + 4,
        )

    if has_any(text, ["paper", "art", "artist"]) or {"book", "scissors"} & cats:
        return (
            "unclear_or_non_creation",
            "creation_vs_salience",
            "none_or_unclear",
            "paper/book/scissors/art object",
            "paper/book/art object",
            "creation-related object is salient but process is not necessarily active",
            "salient materials can appear without a creation process",
            triggers,
            score,
        )
    return None


def download_image(file_name: str, allow_download: bool = True) -> Path | None:
    IMAGE_DIR.mkdir(parents=True, exist_ok=True)
    out_path = IMAGE_DIR / file_name
    if out_path.exists() and out_path.stat().st_size > 0:
        return out_path
    if not allow_download:
        return None
    url = f"http://images.cocodataset.org/val2017/{file_name}"
    for attempt in range(1):
        try:
            with urllib.request.urlopen(url, timeout=12, context=SSL_CONTEXT) as response:
                out_path.write_bytes(response.read())
            if out_path.exists() and out_path.stat().st_size > 0:
                return out_path
        except Exception:
            time.sleep(2**attempt)
    return None


def write_csv(path: Path, rows: list[Candidate]) -> None:
    fieldnames = [
        "id",
        "image_path",
        "keep",
        "question",
        "gold_step_label",
        "neighbor_step_label",
        "boundary_family",
        "why_not_neighbor",
        "visible_actor",
        "visible_tool",
        "visible_target",
        "visible_state_cues",
        "issue_type",
        "notes",
        "source_dataset",
        "source_image_id",
        "suggested_label",
        "suggested_boundary_family",
        "suggested_actor",
        "suggested_tool",
        "suggested_target",
        "suggested_state_cues",
        "suggestion_reason",
        "captions",
        "categories",
        "trigger_words",
        "priority_score",
        "subfamily",
        "scenario_family",
    ]
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()
        for i, row in enumerate(rows, 1):
            writer.writerow(
                {
                    "id": f"task3craft_{i:04d}_{row.image_id}",
                    "image_path": row.image_path,
                    "keep": "",
                    "question": "What step of the creation process is shown in this scene?",
                    "gold_step_label": "",
                    "neighbor_step_label": "",
                    "boundary_family": "",
                    "why_not_neighbor": "",
                    "visible_actor": "",
                    "visible_tool": "",
                    "visible_target": "",
                    "visible_state_cues": "",
                    "issue_type": "",
                    "notes": "",
                    "source_dataset": "COCO val2017",
                    "source_image_id": row.image_id,
                    "suggested_label": row.suggested_label,
                    "suggested_boundary_family": row.suggested_boundary_family,
                    "suggested_actor": row.suggested_actor,
                    "suggested_tool": row.suggested_tool,
                    "suggested_target": row.suggested_target,
                    "suggested_state_cues": row.suggested_state_cues,
                    "suggestion_reason": row.suggestion_reason,
                    "captions": " || ".join(row.captions),
                    "categories": ", ".join(sorted(set(row.categories))),
                    "trigger_words": ", ".join(sorted(set(row.trigger_words))),
                    "priority_score": row.priority_score,
                    "subfamily": row.subfamily,
                    "scenario_family": "craft_art_drawing",
                }
            )


def write_html(path: Path, rows: list[Candidate]) -> None:
    cards = []
    for i, row in enumerate(rows, 1):
        rel = Path(row.image_path)
        try:
            rel = rel.relative_to(path.parent)
        except ValueError:
            pass
        cards.append(
            f"""
            <article class="card">
              <h2>{i:03d} / COCO {row.image_id}</h2>
              <img src="{html.escape(str(rel))}" />
              <p><b>Suggested:</b> {html.escape(row.suggested_label)} / {html.escape(row.suggested_boundary_family)}</p>
              <p><b>Subfamily:</b> {html.escape(row.subfamily)}</p>
              <p><b>Tool-target:</b> {html.escape(row.suggested_tool)} -> {html.escape(row.suggested_target)}</p>
              <p><b>State cues:</b> {html.escape(row.suggested_state_cues)}</p>
              <p><b>Reason:</b> {html.escape(row.suggestion_reason)}</p>
              <p><b>Captions:</b> {html.escape(" || ".join(row.captions))}</p>
              <p><b>Categories:</b> {html.escape(", ".join(sorted(set(row.categories))))}</p>
            </article>
            """
        )
    path.write_text(
        f"""<!doctype html>
<html>
<head>
  <meta charset="utf-8" />
  <title>Task3 P6 Craft / Art / Drawing Review</title>
  <style>
    body {{ font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", sans-serif; margin: 24px; background: #f7f7f5; color: #1d1d1f; }}
    .grid {{ display: grid; grid-template-columns: repeat(auto-fill, minmax(360px, 1fr)); gap: 18px; }}
    .card {{ background: white; border: 1px solid #d9d9d6; border-radius: 8px; padding: 14px; }}
    img {{ width: 100%; max-height: 320px; object-fit: contain; background: #eee; }}
    h1 {{ margin-top: 0; }}
    h2 {{ font-size: 16px; margin: 0 0 10px; }}
    p {{ font-size: 13px; line-height: 1.35; }}
  </style>
</head>
<body>
  <h1>Task3 P6 Craft / Art / Drawing Review</h1>
  <p>Use the xlsx file for annotation. This page is only for visual inspection.</p>
  <div class="grid">
    {''.join(cards)}
  </div>
</body>
</html>
""",
        encoding="utf-8",
    )


def write_xlsx(csv_path: Path, xlsx_path: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    rows = list(csv.DictReader(csv_path.open("r", encoding="utf-8")))
    wb = Workbook()
    ws = wb.active
    ws.title = "craft_review"
    headers = list(rows[0].keys()) if rows else []
    ws.append(headers)
    for row in rows:
        ws.append([row.get(h, "") for h in headers])

    header_fill = PatternFill("solid", fgColor="EDE7DA")
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    widths = {
        "A": 24,
        "B": 58,
        "C": 10,
        "D": 42,
        "E": 24,
        "F": 24,
        "G": 28,
        "H": 42,
        "I": 18,
        "J": 18,
        "K": 22,
        "L": 36,
        "M": 26,
        "N": 32,
        "W": 70,
        "X": 70,
    }
    for col, width in widths.items():
        ws.column_dimensions[col].width = width
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    options = {
        "keep": ["yes", "no"],
        "gold_step_label": ["pre_creation_setup", "active_creation", "active_assembly", "finished_artifact", "unclear_or_non_creation"],
        "neighbor_step_label": ["pre_creation_setup", "active_creation", "active_assembly", "finished_artifact", "unclear_or_non_creation", "none"],
        "boundary_family": ["setup_vs_active", "assembly_vs_finished", "active_vs_finished", "creation_vs_salience"],
        "issue_type": ["ok", "bad_image", "not_creation_process", "label_unclear", "duplicate", "too_hard", "other"],
    }
    header_to_col = {cell.value: cell.column_letter for cell in ws[1]}
    for header, values in options.items():
        col = header_to_col.get(header)
        if not col:
            continue
        dv = DataValidation(type="list", formula1=f'"{",".join(values)}"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"{col}2:{col}{ws.max_row}")

    comments = {
        "keep": "yes means this item is visually suitable and should enter the eval set; no means exclude it.",
        "gold_step_label": "The correct visible creation stage.",
        "neighbor_step_label": "The most tempting but wrong neighboring stage, if there is one.",
        "boundary_family": "The process boundary this item tests.",
        "why_not_neighbor": "One short visual reason why the neighbor label is wrong.",
        "visible_state_cues": "Observable cues, not inferred story.",
    }
    for header, text in comments.items():
        col = header_to_col.get(header)
        if col:
            ws[f"{col}1"].comment = Comment(text, "Codex")

    xlsx_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(xlsx_path)


def build(limit: int, allow_download: bool = True, balanced: bool = True) -> list[Candidate]:
    caps_by_img, cats_by_img, file_by_img = load_coco_metadata()
    metas: list[CandidateMeta] = []
    for image_id, captions in caps_by_img.items():
        categories = cats_by_img.get(image_id, [])
        inferred = infer_candidate(captions, categories)
        if inferred is None:
            continue
        label, boundary, actor, tool, target, state_cues, reason, triggers, score = inferred
        file_name = file_by_img.get(image_id)
        if not file_name:
            continue
        subfamily = infer_subfamily(" ".join(captions).lower(), categories)
        metas.append(
            CandidateMeta(
                image_id=image_id,
                file_name=file_name,
                captions=captions,
                categories=categories,
                suggested_label=label,
                suggested_boundary_family=boundary,
                suggested_actor=actor,
                suggested_tool=tool,
                suggested_target=target,
                suggested_state_cues=state_cues,
                suggestion_reason=reason,
                trigger_words=triggers,
                priority_score=score,
                subfamily=subfamily,
            )
        )
    metas.sort(key=lambda row: (-row.priority_score, row.image_id))

    if balanced:
        quotas = {
            "drawing_painting": 24,
            "writing_paper": 18,
            "cutting_paper": 24,
            "cake_decoration": 18,
            "cake_cutting_or_finished": 18,
            "making_assembly": 18,
            "other_creation_salience": 8,
        }
        selected: list[CandidateMeta] = []
        used: set[int] = set()
        by_family: dict[str, list[CandidateMeta]] = defaultdict(list)
        for meta in metas:
            by_family[meta.subfamily].append(meta)
        for family, quota in quotas.items():
            for meta in by_family.get(family, [])[:quota]:
                if len(selected) >= limit:
                    break
                selected.append(meta)
                used.add(meta.image_id)
        for meta in metas:
            if len(selected) >= limit:
                break
            if meta.image_id not in used:
                selected.append(meta)
                used.add(meta.image_id)
        metas = selected

    rows: list[Candidate] = []
    for meta in metas:
        if len(rows) >= limit:
            break
        image_path = download_image(meta.file_name, allow_download=allow_download)
        if image_path is None:
            continue
        rows.append(
            Candidate(
                image_id=meta.image_id,
                image_path=str(image_path.resolve()),
                captions=meta.captions,
                categories=meta.categories,
                suggested_label=meta.suggested_label,
                suggested_boundary_family=meta.suggested_boundary_family,
                suggested_actor=meta.suggested_actor,
                suggested_tool=meta.suggested_tool,
                suggested_target=meta.suggested_target,
                suggested_state_cues=meta.suggested_state_cues,
                suggestion_reason=meta.suggestion_reason,
                trigger_words=meta.trigger_words,
                priority_score=meta.priority_score,
                subfamily=meta.subfamily,
            )
        )
    return rows


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--limit", type=int, default=120)
    parser.add_argument("--csv", default="review/task3_craft_candidates_v1_review.csv")
    parser.add_argument("--html", default="review/task3_craft_candidates_v1_review.html")
    parser.add_argument("--xlsx", default="review/task3_craft_candidates_v1_review.xlsx")
    parser.add_argument("--summary", default="runs/task3_craft_candidates_v1_summary_2026-07-01.json")
    parser.add_argument("--no-download", action="store_true")
    parser.add_argument("--unbalanced", action="store_true")
    args = parser.parse_args()

    rows = build(args.limit, allow_download=not args.no_download, balanced=not args.unbalanced)
    csv_path = ROOT / args.csv
    html_path = ROOT / args.html
    xlsx_path = ROOT / args.xlsx
    summary_path = ROOT / args.summary
    write_csv(csv_path, rows)
    write_html(html_path, rows)
    write_xlsx(csv_path, xlsx_path)
    summary_path.parent.mkdir(parents=True, exist_ok=True)
    summary_path.write_text(
        json.dumps(
            {
                "n": len(rows),
                "csv": str(csv_path),
                "html": str(html_path),
                "xlsx": str(xlsx_path),
                "labels": {label: sum(1 for r in rows if r.suggested_label == label) for label in sorted({r.suggested_label for r in rows})},
                "boundaries": {
                    family: sum(1 for r in rows if r.suggested_boundary_family == family)
                    for family in sorted({r.suggested_boundary_family for r in rows})
                },
                "subfamilies": {
                    family: sum(1 for r in rows if r.subfamily == family)
                    for family in sorted({r.subfamily for r in rows})
                },
            },
            indent=2,
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )
    print(f"Wrote {len(rows)} candidates")
    print(csv_path)
    print(html_path)
    print(xlsx_path)
    print(summary_path)


if __name__ == "__main__":
    main()
