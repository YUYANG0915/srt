#!/usr/bin/env python3
"""Prepare a focused correction workbook for Task3 P6 craft bootstrap labels."""

from __future__ import annotations

import argparse
import csv
import html
import json
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any


LABEL_OPTIONS = [
    "pre_creation_setup",
    "active_creation",
    "active_assembly",
    "finished_artifact",
    "unclear_or_non_creation",
]
BOUNDARY_OPTIONS = [
    "setup_vs_active",
    "active_vs_finished",
    "assembly_vs_finished",
    "creation_vs_salience",
]
ISSUE_OPTIONS = [
    "ok",
    "bad_image",
    "not_creation_process",
    "label_unclear",
    "duplicate",
    "too_hard",
    "other",
]


def load_jsonl(path: Path) -> list[dict[str, Any]]:
    return [json.loads(line) for line in path.open("r", encoding="utf-8") if line.strip()]


def load_case_table(path: Path) -> list[dict[str, str]]:
    with path.open("r", encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


def majority(values: list[str]) -> str:
    values = [v for v in values if v]
    if not values:
        return ""
    counts = Counter(values)
    best, _ = counts.most_common(1)[0]
    return best


def disagreement_score(gold: str, preds: dict[str, str]) -> int:
    values = [v for v in preds.values() if v]
    model_disagree = sum(1 for v in values if v != gold)
    diversity = len(set(values))
    return model_disagree * 10 + diversity


def build_rows(items_path: Path, case_table_path: Path) -> list[dict[str, Any]]:
    items = {row["id"]: row for row in load_jsonl(items_path)}
    by_id: dict[str, dict[str, str]] = defaultdict(dict)
    for row in load_case_table(case_table_path):
        key = f"{row['model']}::{row['condition']}"
        by_id[row["id"]][key] = row.get("pred", "")

    condition_order = [
        "GPT-5.4-mini::default",
        "GPT-5.4-mini::generic_srt",
        "GPT-5.4-mini::scenario_srt",
        "GPT-5.4-mini::scenario_srt_v2",
        "GPT-5.4-mini::boundary_router",
        "Qwen2.5-VL-72B::default",
        "Qwen2.5-VL-72B::generic_srt",
        "Qwen2.5-VL-72B::scenario_srt",
        "Qwen2.5-VL-72B::scenario_srt_v2",
        "Qwen2.5-VL-72B::boundary_router",
    ]

    rows: list[dict[str, Any]] = []
    for item_id, item in items.items():
        preds = {key: by_id[item_id].get(key, "") for key in condition_order}
        gold = str(item.get("gold_label", ""))
        model_majority = majority(list(preds.values()))
        score = disagreement_score(gold, preds)
        row = {
            "id": item_id,
            "image_path": item.get("image_path", ""),
            "keep": "yes",
            "corrected_gold_step_label": "",
            "corrected_boundary_family": "",
            "corrected_neighbor_step_label": "",
            "why_not_neighbor": item.get("why_not_neighbor", ""),
            "issue_type": "",
            "notes": "",
            "suggested_label": gold,
            "suggested_boundary_family": item.get("boundary_family", ""),
            "model_majority": model_majority,
            "disagreement_score": score,
            "subfamily": item.get("subfamily", ""),
            "visible_tool": item.get("visible_tool", ""),
            "visible_target": item.get("visible_target", ""),
            "visible_state_cues": item.get("visible_state_cues", ""),
            "captions": item.get("captions", ""),
            "categories": item.get("categories", ""),
            "source_image_id": item.get("image_id", ""),
        }
        for key in condition_order:
            row[key.replace("::", "_").replace("-", "").replace(".", "").replace(" ", "")] = preds[key]
        rows.append(row)

    rows.sort(key=lambda row: (-int(row["disagreement_score"]), row["subfamily"], row["id"]))
    return rows


def write_csv(path: Path, rows: list[dict[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    fieldnames = list(rows[0].keys()) if rows else []
    with path.open("w", encoding="utf-8", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)


def write_html(path: Path, rows: list[dict[str, Any]]) -> None:
    cards = []
    for idx, row in enumerate(rows, start=1):
        cards.append(
            f"""
            <article class="card">
              <h2>{idx:03d} / {html.escape(row['id'])} / score {row['disagreement_score']}</h2>
              <img src="{html.escape(row['image_path'])}" />
              <p><b>Suggested:</b> {html.escape(row['suggested_label'])} / {html.escape(row['suggested_boundary_family'])}</p>
              <p><b>Majority:</b> {html.escape(row['model_majority'])}</p>
              <p><b>Subfamily:</b> {html.escape(row['subfamily'])}</p>
              <p><b>Cues:</b> {html.escape(row['visible_tool'])} -> {html.escape(row['visible_target'])}; {html.escape(row['visible_state_cues'])}</p>
              <p><b>Captions:</b> {html.escape(row['captions'])}</p>
            </article>
            """
        )
    path.write_text(
        f"""<!doctype html>
<html>
<head>
  <meta charset="utf-8" />
  <title>Task3 P6 Craft Focused Correction</title>
  <style>
    body {{ font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", sans-serif; margin: 24px; background: #f7f7f5; color: #1d1d1f; }}
    .grid {{ display: grid; grid-template-columns: repeat(auto-fill, minmax(380px, 1fr)); gap: 18px; }}
    .card {{ background: white; border: 1px solid #d9d9d6; border-radius: 8px; padding: 14px; }}
    img {{ width: 100%; max-height: 320px; object-fit: contain; background: #eee; }}
    h1 {{ margin-top: 0; }}
    h2 {{ font-size: 16px; margin: 0 0 10px; }}
    p {{ font-size: 13px; line-height: 1.35; }}
  </style>
</head>
<body>
  <h1>Task3 P6 Craft Focused Correction</h1>
  <p>Rows are sorted by model disagreement. Use the xlsx for correction.</p>
  <div class="grid">{''.join(cards)}</div>
</body>
</html>
""",
        encoding="utf-8",
    )


def write_xlsx(csv_path: Path, xlsx_path: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    with csv_path.open("r", encoding="utf-8-sig", newline="") as f:
        rows = list(csv.DictReader(f))
    wb = Workbook()
    ws = wb.active
    ws.title = "craft_correction"
    headers = list(rows[0].keys()) if rows else []
    ws.append(headers)
    for row in rows:
        ws.append([row.get(h, "") for h in headers])

    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="EDE7DA")
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    widths = {
        "A": 24,
        "B": 58,
        "C": 10,
        "D": 26,
        "E": 26,
        "F": 26,
        "G": 42,
        "H": 22,
        "I": 32,
        "J": 24,
        "K": 26,
        "L": 24,
        "M": 16,
        "N": 22,
        "O": 22,
        "P": 26,
        "Q": 42,
        "R": 72,
    }
    for col, width in widths.items():
        ws.column_dimensions[col].width = width
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    header_to_col = {cell.value: cell.column_letter for cell in ws[1]}
    options = {
        "keep": ["yes", "no"],
        "corrected_gold_step_label": LABEL_OPTIONS,
        "corrected_boundary_family": BOUNDARY_OPTIONS,
        "corrected_neighbor_step_label": LABEL_OPTIONS + ["none"],
        "issue_type": ISSUE_OPTIONS,
    }
    for header, values in options.items():
        col = header_to_col.get(header)
        if not col:
            continue
        dv = DataValidation(type="list", formula1=f'"{",".join(values)}"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"{col}2:{col}{ws.max_row}")

    comments = {
        "corrected_gold_step_label": "Fill this when suggested_label is wrong or uncertain. Leave blank to accept suggested_label.",
        "corrected_boundary_family": "Fill this when suggested_boundary_family is wrong. Leave blank to accept suggested_boundary_family.",
        "corrected_neighbor_step_label": "Tempting wrong neighboring stage; optional.",
        "why_not_neighbor": "Short observable reason why the neighbor label is wrong; optional but useful.",
        "issue_type": "Use not_creation_process/bad_image/too_hard when keep should be no.",
    }
    for header, text in comments.items():
        col = header_to_col.get(header)
        if col:
            ws[f"{col}1"].comment = Comment(text, "Codex")

    xlsx_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(xlsx_path)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--items", default="data/task3_craft_eval_v1_reviewed.jsonl")
    parser.add_argument("--case-table", default="runs/task3_craft_v1bootstrap79_gpt_qwen_5cond_summary_2026-07-01/task3_craft_case_table.csv")
    parser.add_argument("--csv", default="review/task3_craft_focused_correction_v1.csv")
    parser.add_argument("--html", default="review/task3_craft_focused_correction_v1.html")
    parser.add_argument("--xlsx", default="review/task3_craft_focused_correction_v1.xlsx")
    args = parser.parse_args()

    rows = build_rows(Path(args.items), Path(args.case_table))
    write_csv(Path(args.csv), rows)
    write_html(Path(args.html), rows)
    write_xlsx(Path(args.csv), Path(args.xlsx))
    print(f"Wrote {len(rows)} focused correction rows")
    print(Path(args.csv).resolve())
    print(Path(args.html).resolve())
    print(Path(args.xlsx).resolve())


if __name__ == "__main__":
    main()
