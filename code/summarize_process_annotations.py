#!/usr/bin/env python3
"""Clean and summarize manual process-failure annotations."""

from __future__ import annotations

import json
from collections import Counter
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.worksheet.datavalidation import DataValidation


BASE = Path(__file__).resolve().parents[1]
IN_XLSX = BASE / "runs/process_failure_annotation/process_failure_annotation_template.xlsx"
OUT_DIR = BASE / "runs/process_failure_annotation"
HUMAN_COLS = [
    "human_keep",
    "human_failure_label_primary",
    "human_failure_label_secondary",
    "human_failure_source",
    "human_process_answer_consistent",
    "human_expected_policy",
    "human_notes",
]


def norm(value):
    if pd.isna(value):
        return ""
    return str(value).strip()


def has_any_annotation(row: pd.Series) -> bool:
    return any(norm(row[col]) for col in HUMAN_COLS)


def apply_defaults(row: pd.Series) -> pd.Series:
    keep = norm(row["human_keep"])
    primary = norm(row["human_failure_label_primary"])
    case_type = norm(row["case_type"])
    task = norm(row["task"])

    source_defaults = {
        "candidate_detection_failure": "perception",
        "candidate_relevance_failure": "grounding",
        "uniqueness_judgment_failure": "uniqueness_judgment",
        "process_answer_inconsistency": "pragmatic_policy",
        "salience_guessing": "pragmatic_policy",
        "over_conservative_clarification": "pragmatic_policy",
        "none_or_repaired": "none",
        "none_control_success": "none",
        "invalid_item": "other",
    }
    consistency_defaults = {
        "process_answer_inconsistency": "no",
        "salience_guessing": "no",
        "none_or_repaired": "yes",
        "none_control_success": "yes",
    }

    if keep == "no":
        row["human_failure_label_primary"] = primary or "invalid_item"
        row["human_failure_source"] = norm(row["human_failure_source"]) or "other"
        row["human_process_answer_consistent"] = norm(row["human_process_answer_consistent"]) or "unclear"
        row["human_expected_policy"] = norm(row["human_expected_policy"]) or "uncertain"
        return row

    if keep == "yes" and not primary:
        if case_type == "T2_unique_control_visual_answers":
            row["human_failure_label_primary"] = "none_control_success"
            row["human_failure_source"] = norm(row["human_failure_source"]) or "none"
            row["human_process_answer_consistent"] = norm(row["human_process_answer_consistent"]) or "yes"
            row["human_expected_policy"] = norm(row["human_expected_policy"]) or "answer"
        elif task == "task1" and norm(row["visual_policy"]) in {"clarify", "enumerate"}:
            row["human_failure_label_primary"] = "none_or_repaired"
            row["human_failure_source"] = norm(row["human_failure_source"]) or "none"
            row["human_process_answer_consistent"] = norm(row["human_process_answer_consistent"]) or "yes"
            row["human_expected_policy"] = norm(row["human_expected_policy"]) or norm(row["visual_policy"])
    primary = norm(row["human_failure_label_primary"])
    if keep == "yes" and primary:
        row["human_failure_source"] = norm(row["human_failure_source"]) or source_defaults.get(primary, "")
        row["human_process_answer_consistent"] = (
            norm(row["human_process_answer_consistent"]) or consistency_defaults.get(primary, "unclear")
        )
        if not norm(row["human_expected_policy"]):
            row["human_expected_policy"] = "answer" if task == "task2" else "clarify"
    return row


def propagate_repeated_rows(df: pd.DataFrame) -> pd.DataFrame:
    out = df.copy()
    for col in HUMAN_COLS:
        out[col] = out[col].astype("object")
    key_cols = ["case_type", "task", "sample_id"]
    annotated = out[out.apply(has_any_annotation, axis=1)].copy()
    lookup = {}
    for _, row in annotated.iterrows():
        key = tuple(norm(row[col]) for col in key_cols)
        if key not in lookup:
            lookup[key] = {col: row[col] for col in HUMAN_COLS}

    for idx, row in out.iterrows():
        if has_any_annotation(row):
            continue
        key = tuple(norm(row[col]) for col in key_cols)
        if key in lookup:
            for col, value in lookup[key].items():
                out.at[idx, col] = value
            note = norm(out.at[idx, "human_notes"])
            out.at[idx, "human_notes"] = note or "propagated_from_duplicate_sample"
    return out


def write_jsonl(df: pd.DataFrame, path: Path) -> None:
    with path.open("w", encoding="utf-8") as f:
        for record in df.fillna("").to_dict(orient="records"):
            f.write(json.dumps(record, ensure_ascii=False) + "\n")


def add_validations(xlsx_path: Path) -> None:
    wb = load_workbook(xlsx_path)
    ws = wb["Annotations"]
    validations = {
        "AC": ["yes", "no"],
        "AD": [
            "candidate_detection_failure",
            "candidate_relevance_failure",
            "uniqueness_judgment_failure",
            "process_answer_inconsistency",
            "salience_guessing",
            "over_conservative_clarification",
            "none_or_repaired",
            "none_control_success",
            "invalid_item",
            "other",
        ],
        "AE": [
            "candidate_detection_failure",
            "candidate_relevance_failure",
            "uniqueness_judgment_failure",
            "process_answer_inconsistency",
            "salience_guessing",
            "over_conservative_clarification",
            "none_or_repaired",
            "none_control_success",
            "invalid_item",
            "other",
        ],
        "AF": ["perception", "grounding", "uniqueness_judgment", "pragmatic_policy", "none", "other"],
        "AG": ["yes", "no", "unclear"],
        "AH": ["answer", "clarify", "enumerate", "uncertain"],
    }
    for col, options in validations.items():
        dv = DataValidation(type="list", formula1=f'"{",".join(options)}"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"{col}2:{col}{ws.max_row}")
    wb.save(xlsx_path)


def summarize(df: pd.DataFrame, name: str) -> list[str]:
    rows = []
    keep_yes = df[df["human_keep"].fillna("").astype(str).str.strip() == "yes"]
    kept = len(keep_yes)
    rows += [f"## {name}", "", f"- Rows: {len(df)}", f"- Kept rows: {kept}", ""]

    for title, col in [
        ("Primary failure labels", "human_failure_label_primary"),
        ("Failure sources", "human_failure_source"),
        ("Expected policy", "human_expected_policy"),
        ("Process-answer consistency", "human_process_answer_consistent"),
    ]:
        counts = Counter(norm(v) or "<blank>" for v in keep_yes[col])
        rows += [f"### {title}", "", "| value | count |", "|---|---:|"]
        for key, value in counts.most_common():
            rows.append(f"| `{key}` | {value} |")
        rows.append("")

    rows += ["### By task/model", "", "| task | model | kept |", "|---|---|---:|"]
    grouped = keep_yes.groupby(["task", "model"], dropna=False).size().reset_index(name="n")
    for _, row in grouped.iterrows():
        rows.append(f"| `{row['task']}` | `{row['model']}` | {row['n']} |")
    rows.append("")
    return rows


def main() -> None:
    df = pd.read_excel(IN_XLSX, sheet_name="Annotation")
    for col in HUMAN_COLS:
        if col not in df:
            df[col] = ""

    strict = df[df.apply(has_any_annotation, axis=1)].copy()
    strict = strict.apply(apply_defaults, axis=1)

    propagated = propagate_repeated_rows(df)
    propagated = propagated[propagated.apply(has_any_annotation, axis=1)].copy()
    propagated = propagated.apply(apply_defaults, axis=1)

    outputs = [
        ("strict", strict, OUT_DIR / "process_failure_annotations_strict"),
        ("propagated", propagated, OUT_DIR / "process_failure_annotations_propagated"),
    ]
    for _, frame, stem in outputs:
        frame.to_csv(stem.with_suffix(".csv"), index=False)
        write_jsonl(frame, stem.with_suffix(".jsonl"))
        with pd.ExcelWriter(stem.with_suffix(".xlsx"), engine="openpyxl") as writer:
            frame.to_excel(writer, sheet_name="Annotations", index=False)
        add_validations(stem.with_suffix(".xlsx"))

    md = [
        "# Process Failure Annotation Summary",
        "",
        "Two versions are reported:",
        "- `strict`: only rows with manual annotation in the workbook, with obvious defaults filled for kept Task2 controls and invalid rows.",
        "- `propagated`: copies manual annotation to blank rows with the same `(case_type, task, sample_id)`, matching the user's note that blank lower rows duplicate upper rows.",
        "",
    ]
    md.extend(summarize(strict, "Strict"))
    md.extend(summarize(propagated, "Propagated"))
    (OUT_DIR / "annotation_summary.md").write_text("\n".join(md), encoding="utf-8")
    print(OUT_DIR / "annotation_summary.md")
    print(f"strict rows={len(strict)} propagated rows={len(propagated)}")


if __name__ == "__main__":
    main()
