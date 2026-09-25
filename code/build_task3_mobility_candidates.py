#!/usr/bin/env python3
"""Build Task3 P7 mobility / transport process candidate review pack from COCO val2017."""

from __future__ import annotations

import argparse
import csv
import html
import io
import json
import re
import ssl
import time
import urllib.request
import zipfile
from collections import defaultdict
from dataclasses import dataclass
from pathlib import Path

try:
    import certifi
except ModuleNotFoundError:
    certifi = None


ROOT = Path(__file__).resolve().parents[1]
DATA_DIR = ROOT / "data"
ZIP_PATH = DATA_DIR / "coco_annotations" / "annotations_trainval2017.zip"
IMAGE_DIR = DATA_DIR / "task3_mobility_coco_images"
SSL_CONTEXT = ssl.create_default_context(cafile=certifi.where()) if certifi else ssl.create_default_context()


@dataclass
class Candidate:
    image_id: int
    image_path: str
    captions: list[str]
    categories: list[str]
    suggested_label: str
    suggested_boundary_family: str
    suggested_actor: str
    suggested_vehicle: str
    suggested_relation: str
    suggested_state_cues: str
    suggestion_reason: str
    trigger_words: list[str]
    priority_score: int
    subfamily: str


@dataclass
class CandidateMeta:
    image_id: int
    file_name: str
    captions: list[str]
    categories: list[str]
    suggested_label: str
    suggested_boundary_family: str
    suggested_actor: str
    suggested_vehicle: str
    suggested_relation: str
    suggested_state_cues: str
    suggestion_reason: str
    trigger_words: list[str]
    priority_score: int
    subfamily: str


def load_coco_metadata() -> tuple[dict[int, list[str]], dict[int, list[str]], dict[int, str]]:
    with zipfile.ZipFile(ZIP_PATH) as zf:
        caps = json.load(io.TextIOWrapper(zf.open("annotations/captions_val2017.json"), encoding="utf-8"))
        inst = json.load(io.TextIOWrapper(zf.open("annotations/instances_val2017.json"), encoding="utf-8"))

    cat_name = {row["id"]: row["name"] for row in inst["categories"]}
    caps_by_img: dict[int, list[str]] = defaultdict(list)
    cats_by_img: dict[int, list[str]] = defaultdict(list)
    file_by_img: dict[int, str] = {}
    for image in inst["images"]:
        file_by_img[image["id"]] = image["file_name"]
    for ann in caps["annotations"]:
        caps_by_img[ann["image_id"]].append(ann["caption"])
    for ann in inst["annotations"]:
        cats_by_img[ann["image_id"]].append(cat_name[ann["category_id"]])
    return caps_by_img, cats_by_img, file_by_img


def contains_any(text: str, terms: list[str]) -> list[str]:
    return [term for term in terms if term in text]


def has_any(text: str, terms: list[str]) -> bool:
    return any(term in text for term in terms)


def infer_subfamily(text: str, categories: list[str]) -> str:
    cats = set(categories)
    if "bicycle" in cats or has_any(text, ["bicycle", "bike", "bicycl"]):
        return "bicycle"
    if "motorcycle" in cats or has_any(text, ["motorcycle", "motor bike", "motorbike", "moped"]):
        return "motorcycle"
    if "skateboard" in cats or has_any(text, ["skateboard", "skate board"]):
        return "skateboard"
    if "horse" in cats or "horse" in text:
        return "horse_riding"
    if {"skis", "snowboard"} & cats or has_any(text, ["skiing", "snowboard", "snowboarding"]):
        return "snow_sport_mobility"
    if {"surfboard", "boat"} & cats or has_any(text, ["surfing", "surfboard", "boat", "kayak", "canoe"]):
        return "water_mobility"
    if {"car", "bus", "truck", "train", "airplane"} & cats:
        return "vehicle_traffic"
    return "other_mobility"


def infer_vehicle(text: str, categories: list[str]) -> str:
    ordered = [
        "bicycle",
        "motorcycle",
        "skateboard",
        "horse",
        "skis",
        "snowboard",
        "surfboard",
        "boat",
        "car",
        "bus",
        "truck",
        "train",
        "airplane",
    ]
    cats = set(categories)
    for cat in ordered:
        if cat in cats:
            return cat
    for word in ["bike", "bicycle", "motorcycle", "skateboard", "horse", "ski", "snowboard", "surfboard", "boat", "car", "bus", "train"]:
        if word in text:
            return word
    return "vehicle/mobility device"


def infer_candidate(captions: list[str], categories: list[str]) -> tuple[str, str, str, str, str, str, str, list[str], int] | None:
    text = " ".join(captions).lower()
    cats = set(categories)
    trigger_terms = [
        "riding",
        "rides",
        "ride",
        "driving",
        "drives",
        "driver",
        "parked",
        "parking",
        "bicycle",
        "bike",
        "motorcycle",
        "skateboard",
        "skiing",
        "snowboard",
        "surfing",
        "boat",
        "train",
        "bus",
        "car",
        "truck",
        "airplane",
        "horse",
        "sitting on",
        "standing next to",
        "next to a bike",
        "bus stop",
        "waiting for the bus",
        "waiting to board",
    ]
    vehicle_cats = {
        "bicycle",
        "motorcycle",
        "skateboard",
        "skis",
        "snowboard",
        "surfboard",
        "boat",
        "car",
        "bus",
        "truck",
        "train",
        "airplane",
        "horse",
    }
    triggers = contains_any(text, trigger_terms)
    triggers.extend(sorted(vehicle_cats & cats))
    if not triggers:
        return None
    vehicle_text_terms = [
        "bike",
        "bicycle",
        "motorcycle",
        "moped",
        "skateboard",
        "ski",
        "snowboard",
        "surfboard",
        "surfing",
        "boat",
        "kayak",
        "canoe",
        "car",
        "bus",
        "truck",
        "train",
        "airplane",
        "plane",
        "horse",
        "bus stop",
        "waiting for the bus",
        "waiting to board",
    ]
    if not (vehicle_cats & cats or has_any(text, vehicle_text_terms)):
        return None

    vehicle = infer_vehicle(text, categories)
    score = len(set(triggers))
    if "person" in cats:
        score += 3
    if has_any(text, ["riding", "rides", "driving", "skiing", "snowboarding", "surfing", "skateboarding", "flying", "flies", "landing", "take off", "taking off"]):
        score += 8
    if has_any(text, ["parked", "parking", "sitting parked", "parked on"]):
        score += 5
    if {"bicycle", "motorcycle", "skateboard", "horse", "skis", "snowboard", "surfboard"} & cats:
        score += 3

    if has_any(text, ["parked", "parking", "sitting parked", "parked on", "parked next to"]):
        return (
            "parked_or_finished",
            "active_vs_parked",
            "none_or_unclear",
            vehicle,
            "vehicle is stationary/parked",
            "vehicle appears as a stopped object rather than active mobility",
            "caption describes a parked/stationary result state",
            triggers,
            score + 3,
        )

    if has_any(text, ["riding", "rides", "driving", "skiing", "snowboarding", "surfing", "skateboarding", "pulling a cart", "on a horse", "flying", "flies", "landing", "taking off", "take off"]):
        return (
            "active_mobility",
            "setup_vs_active",
            "person/rider",
            vehicle,
            "person is on/in/control of mobility device",
            "rider/driver posture and contact suggest active movement or mobility",
            "caption directly describes active mobility",
            triggers,
            score + 7,
        )

    setup_patterns = [
        r"sitting on (a |the |their )?(bike|bicycle|motorcycle|moped|skateboard|horse)",
        r"standing (next to|near|beside|around|in front of) (a |the |their |his |her )?(bike|bicycle|motorcycle|moped|skateboard|horse|bus|train)",
        r"(next to|near|beside) (a |the |their |his |her )?(bike|bicycle|motorcycle|moped|skateboard|horse|bus|train)",
        r"holding (a |the |his |her )?(bike|bicycle|skateboard)",
        r"waiting (for|at|to board)",
        r"bus stop",
    ]
    if any(re.search(pattern, text) for pattern in setup_patterns):
        return (
            "pre_mobility_setup",
            "setup_vs_active",
            "person",
            vehicle,
            "near/on vehicle but active motion is unclear",
            "person-vehicle proximity suggests setup or staging before mobility",
            "caption has setup-like person-vehicle relation without clear active motion",
            triggers,
            score + 3,
        )

    if {"car", "bus", "truck", "train", "airplane", "boat"} & cats and not has_any(
        text,
        ["riding", "driving", "parked", "parking", "flying", "flies", "landing", "take off", "taking off"],
    ):
        return (
            "unclear_or_non_mobility",
            "mobility_vs_salience",
            "none_or_unclear",
            vehicle,
            "vehicle is salient but process stage is not grounded",
            "vehicle presence alone does not identify setup, active movement, or parked/result stage",
            "caption/category has vehicle salience without a clear process boundary",
            triggers,
            score,
        )
    return None


def download_image(file_name: str, allow_download: bool = True) -> Path | None:
    IMAGE_DIR.mkdir(parents=True, exist_ok=True)
    out_path = IMAGE_DIR / file_name
    if out_path.exists() and out_path.stat().st_size > 0:
        return out_path
    if not allow_download:
        return None
    url = f"http://images.cocodataset.org/val2017/{file_name}"
    for attempt in range(2):
        try:
            with urllib.request.urlopen(url, timeout=20, context=SSL_CONTEXT) as response:
                out_path.write_bytes(response.read())
            if out_path.exists() and out_path.stat().st_size > 0:
                return out_path
        except Exception:
            time.sleep(2**attempt)
    return None


def write_csv(path: Path, rows: list[Candidate]) -> None:
    fieldnames = [
        "id",
        "image_path",
        "keep",
        "question",
        "gold_step_label",
        "neighbor_step_label",
        "boundary_family",
        "why_not_neighbor",
        "visible_actor",
        "visible_vehicle",
        "visible_relation",
        "visible_state_cues",
        "issue_type",
        "notes",
        "source_dataset",
        "source_image_id",
        "suggested_label",
        "suggested_boundary_family",
        "suggested_actor",
        "suggested_vehicle",
        "suggested_relation",
        "suggested_state_cues",
        "suggestion_reason",
        "captions",
        "categories",
        "trigger_words",
        "priority_score",
        "subfamily",
        "scenario_family",
    ]
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()
        for i, row in enumerate(rows, 1):
            writer.writerow(
                {
                    "id": f"task3mobile_{i:04d}_{row.image_id}",
                    "image_path": row.image_path,
                    "keep": "",
                    "question": "What stage of the mobility process is shown in this scene?",
                    "gold_step_label": "",
                    "neighbor_step_label": "",
                    "boundary_family": "",
                    "why_not_neighbor": "",
                    "visible_actor": "",
                    "visible_vehicle": "",
                    "visible_relation": "",
                    "visible_state_cues": "",
                    "issue_type": "",
                    "notes": "",
                    "source_dataset": "COCO val2017",
                    "source_image_id": row.image_id,
                    "suggested_label": row.suggested_label,
                    "suggested_boundary_family": row.suggested_boundary_family,
                    "suggested_actor": row.suggested_actor,
                    "suggested_vehicle": row.suggested_vehicle,
                    "suggested_relation": row.suggested_relation,
                    "suggested_state_cues": row.suggested_state_cues,
                    "suggestion_reason": row.suggestion_reason,
                    "captions": " || ".join(row.captions),
                    "categories": ", ".join(sorted(set(row.categories))),
                    "trigger_words": ", ".join(sorted(set(row.trigger_words))),
                    "priority_score": row.priority_score,
                    "subfamily": row.subfamily,
                    "scenario_family": "mobility_transport_process",
                }
            )


def write_html(path: Path, rows: list[Candidate]) -> None:
    cards = []
    for i, row in enumerate(rows, 1):
        rel = Path(row.image_path)
        try:
            rel = rel.relative_to(path.parent)
        except ValueError:
            pass
        cards.append(
            f"""
            <article class="card">
              <h2>{i:03d} / COCO {row.image_id}</h2>
              <img src="{html.escape(str(rel))}" />
              <p><b>Suggested:</b> {html.escape(row.suggested_label)} / {html.escape(row.suggested_boundary_family)}</p>
              <p><b>Subfamily:</b> {html.escape(row.subfamily)}</p>
              <p><b>Actor-vehicle:</b> {html.escape(row.suggested_actor)} -> {html.escape(row.suggested_vehicle)}</p>
              <p><b>Relation:</b> {html.escape(row.suggested_relation)}</p>
              <p><b>State cues:</b> {html.escape(row.suggested_state_cues)}</p>
              <p><b>Reason:</b> {html.escape(row.suggestion_reason)}</p>
              <p><b>Captions:</b> {html.escape(" || ".join(row.captions))}</p>
              <p><b>Categories:</b> {html.escape(", ".join(sorted(set(row.categories))))}</p>
            </article>
            """
        )
    path.write_text(
        f"""<!doctype html>
<html>
<head>
  <meta charset="utf-8" />
  <title>Task3 P7 Mobility / Transport Review</title>
  <style>
    body {{ font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", sans-serif; margin: 24px; background: #f7f7f5; color: #1d1d1f; }}
    .grid {{ display: grid; grid-template-columns: repeat(auto-fill, minmax(360px, 1fr)); gap: 18px; }}
    .card {{ background: white; border: 1px solid #d9d9d6; border-radius: 8px; padding: 14px; }}
    img {{ width: 100%; max-height: 320px; object-fit: contain; background: #eee; }}
    h1 {{ margin-top: 0; }}
    h2 {{ font-size: 16px; margin: 0 0 10px; }}
    p {{ font-size: 13px; line-height: 1.35; }}
  </style>
</head>
<body>
  <h1>Task3 P7 Mobility / Transport Review</h1>
  <p>Use the xlsx file for annotation. This page is only for visual inspection.</p>
  <div class="grid">
    {''.join(cards)}
  </div>
</body>
</html>
""",
        encoding="utf-8",
    )


def write_xlsx(csv_path: Path, xlsx_path: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    rows = list(csv.DictReader(csv_path.open("r", encoding="utf-8")))
    wb = Workbook()
    ws = wb.active
    ws.title = "mobility_review"
    headers = list(rows[0].keys()) if rows else []
    ws.append(headers)
    for row in rows:
        ws.append([row.get(h, "") for h in headers])

    header_fill = PatternFill("solid", fgColor="DEEAF6")
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    widths = {
        "A": 24,
        "B": 58,
        "C": 10,
        "D": 42,
        "E": 24,
        "F": 24,
        "G": 28,
        "H": 42,
        "I": 18,
        "J": 20,
        "K": 34,
        "L": 36,
        "M": 26,
        "N": 32,
        "W": 70,
        "X": 70,
    }
    for col, width in widths.items():
        ws.column_dimensions[col].width = width
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    options = {
        "keep": ["yes", "no"],
        "gold_step_label": ["pre_mobility_setup", "active_mobility", "parked_or_finished", "unclear_or_non_mobility"],
        "neighbor_step_label": ["pre_mobility_setup", "active_mobility", "parked_or_finished", "unclear_or_non_mobility", "none"],
        "boundary_family": ["setup_vs_active", "active_vs_parked", "mobility_vs_salience"],
        "issue_type": ["ok", "bad_image", "not_mobility_process", "label_unclear", "duplicate", "too_hard", "other"],
    }
    header_to_col = {cell.value: cell.column_letter for cell in ws[1]}
    for header, values in options.items():
        col = header_to_col.get(header)
        if not col:
            continue
        dv = DataValidation(type="list", formula1=f'"{",".join(values)}"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"{col}2:{col}{ws.max_row}")

    comments = {
        "keep": "yes means this item is visually suitable and should enter the eval set; no means exclude it.",
        "gold_step_label": "The correct visible mobility stage.",
        "neighbor_step_label": "The tempting but wrong neighboring stage.",
        "boundary_family": "The process boundary this item tests.",
        "why_not_neighbor": "One short visual reason why the neighbor label is wrong.",
        "visible_state_cues": "Observable cues only, not inferred story.",
    }
    for header, text in comments.items():
        col = header_to_col.get(header)
        if col:
            ws[f"{col}1"].comment = Comment(text, "Codex")

    xlsx_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(xlsx_path)


def build(limit: int, allow_download: bool = True, balanced: bool = True) -> list[Candidate]:
    caps_by_img, cats_by_img, file_by_img = load_coco_metadata()
    metas: list[CandidateMeta] = []
    for image_id, captions in caps_by_img.items():
        categories = cats_by_img.get(image_id, [])
        inferred = infer_candidate(captions, categories)
        if inferred is None:
            continue
        label, boundary, actor, vehicle, relation, state_cues, reason, triggers, score = inferred
        file_name = file_by_img.get(image_id)
        if not file_name:
            continue
        metas.append(
            CandidateMeta(
                image_id=image_id,
                file_name=file_name,
                captions=captions,
                categories=categories,
                suggested_label=label,
                suggested_boundary_family=boundary,
                suggested_actor=actor,
                suggested_vehicle=vehicle,
                suggested_relation=relation,
                suggested_state_cues=state_cues,
                suggestion_reason=reason,
                trigger_words=triggers,
                priority_score=score,
                subfamily=infer_subfamily(" ".join(captions).lower(), categories),
            )
        )

    metas.sort(key=lambda row: row.priority_score, reverse=True)
    selected: list[CandidateMeta] = []
    if balanced:
        labels = ["active_mobility", "parked_or_finished", "pre_mobility_setup", "unclear_or_non_mobility"]
        by_label_subfamily: dict[tuple[str, str], list[CandidateMeta]] = defaultdict(list)
        by_label: dict[str, list[CandidateMeta]] = defaultdict(list)
        for meta in metas:
            by_label_subfamily[(meta.suggested_label, meta.subfamily)].append(meta)
            by_label[meta.suggested_label].append(meta)
        per_label = max(1, limit // len(labels))
        for label in labels:
            label_rows: list[CandidateMeta] = []
            subfamilies = sorted({sub for lab, sub in by_label_subfamily if lab == label})
            per_subfamily = max(2, per_label // max(1, len(subfamilies)))
            for subfamily in subfamilies:
                label_rows.extend(by_label_subfamily[(label, subfamily)][:per_subfamily])
            if len(label_rows) < per_label:
                seen = {meta.image_id for meta in label_rows}
                label_rows.extend([meta for meta in by_label[label] if meta.image_id not in seen][: per_label - len(label_rows)])
            selected.extend(label_rows[:per_label])
        if len(selected) < limit:
            seen = {meta.image_id for meta in selected}
            selected.extend([meta for meta in metas if meta.image_id not in seen][: limit - len(selected)])
        selected = selected[:limit]
    else:
        selected = metas[:limit]

    rows: list[Candidate] = []
    for meta in selected:
        image_path = download_image(meta.file_name, allow_download=allow_download)
        if image_path is None:
            continue
        rows.append(
            Candidate(
                image_id=meta.image_id,
                image_path=str(image_path.resolve()),
                captions=meta.captions,
                categories=meta.categories,
                suggested_label=meta.suggested_label,
                suggested_boundary_family=meta.suggested_boundary_family,
                suggested_actor=meta.suggested_actor,
                suggested_vehicle=meta.suggested_vehicle,
                suggested_relation=meta.suggested_relation,
                suggested_state_cues=meta.suggested_state_cues,
                suggestion_reason=meta.suggestion_reason,
                trigger_words=meta.trigger_words,
                priority_score=meta.priority_score,
                subfamily=meta.subfamily,
            )
        )
    return rows


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--limit", type=int, default=120)
    parser.add_argument("--csv", default="review/task3_mobility_candidates_v1_review.csv")
    parser.add_argument("--html", default="review/task3_mobility_candidates_v1_review.html")
    parser.add_argument("--xlsx", default="review/task3_mobility_candidates_v1_review.xlsx")
    parser.add_argument("--summary", default="runs/task3_mobility_candidates_v1_summary_2026-07-02.json")
    parser.add_argument("--no-download", action="store_true")
    parser.add_argument("--unbalanced", action="store_true")
    args = parser.parse_args()

    rows = build(args.limit, allow_download=not args.no_download, balanced=not args.unbalanced)
    write_csv(Path(args.csv), rows)
    write_html(Path(args.html), rows)
    write_xlsx(Path(args.csv), Path(args.xlsx))

    summary = {
        "n": len(rows),
        "csv": str(Path(args.csv).resolve()),
        "html": str(Path(args.html).resolve()),
        "xlsx": str(Path(args.xlsx).resolve()),
        "image_dir": str(IMAGE_DIR.resolve()),
        "label_distribution": {},
        "boundary_distribution": {},
        "subfamily_distribution": {},
    }
    for row in rows:
        summary["label_distribution"][row.suggested_label] = summary["label_distribution"].get(row.suggested_label, 0) + 1
        summary["boundary_distribution"][row.suggested_boundary_family] = summary["boundary_distribution"].get(row.suggested_boundary_family, 0) + 1
        summary["subfamily_distribution"][row.subfamily] = summary["subfamily_distribution"].get(row.subfamily, 0) + 1
    Path(args.summary).parent.mkdir(parents=True, exist_ok=True)
    Path(args.summary).write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
