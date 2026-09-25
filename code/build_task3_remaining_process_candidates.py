#!/usr/bin/env python3
"""Build review packs for the remaining SRT process families from COCO val2017.

This script is intentionally conservative: it creates human-reviewable candidate
sets with prefilled labels and boundary families, rather than treating caption
heuristics as final gold labels.
"""

from __future__ import annotations

import argparse
import csv
import html
import io
import json
import re
import ssl
import time
import urllib.request
import zipfile
from collections import defaultdict
from dataclasses import dataclass
from pathlib import Path

try:
    import certifi
except ModuleNotFoundError:
    certifi = None


ROOT = Path(__file__).resolve().parents[1]
DATA_DIR = ROOT / "data"
ZIP_PATH = DATA_DIR / "coco_annotations" / "annotations_trainval2017.zip"
SSL_CONTEXT = ssl.create_default_context(cafile=certifi.where()) if certifi else ssl.create_default_context()


@dataclass(frozen=True)
class FamilyConfig:
    key: str
    title: str
    scenario_family: str
    question: str
    labels: list[str]
    boundaries: list[str]


@dataclass
class Candidate:
    family: FamilyConfig
    image_id: int
    image_path: str
    captions: list[str]
    categories: list[str]
    suggested_label: str
    suggested_boundary_family: str
    suggested_actor: str
    suggested_object: str
    suggested_relation: str
    suggested_state_cues: str
    suggestion_reason: str
    trigger_words: list[str]
    priority_score: int
    subfamily: str


FAMILIES: dict[str, FamilyConfig] = {
    "assembly": FamilyConfig(
        key="assembly",
        title="Assembly / Construction",
        scenario_family="assembly_construction_process",
        question="What stage of the assembly or construction process is shown in this scene?",
        labels=["pre_assembly_setup", "active_assembly", "assembled_or_finished", "unclear_or_non_process"],
        boundaries=["setup_vs_active", "assembly_vs_finished", "assembly_vs_salience"],
    ),
    "physical": FamilyConfig(
        key="physical",
        title="Physical State Transition",
        scenario_family="physical_state_transition_process",
        question="What stage of the physical state transition is shown in this scene?",
        labels=["pre_transition_stable", "active_transition", "post_transition_result", "unclear_or_non_process"],
        boundaries=["stable_vs_active", "active_vs_result", "transition_vs_salience"],
    ),
    "social": FamilyConfig(
        key="social",
        title="Social / Event Script",
        scenario_family="social_event_script_process",
        question="What stage of the social event process is shown in this scene?",
        labels=["pre_event_setup", "active_event", "post_event_result", "unclear_or_non_process"],
        boundaries=["setup_vs_active_event", "active_vs_after_event", "event_vs_salience"],
    ),
    "traffic": FamilyConfig(
        key="traffic",
        title="Navigation / Traffic Interaction",
        scenario_family="navigation_traffic_process",
        question="What stage of the navigation or traffic interaction is shown in this scene?",
        labels=["pre_navigation_waiting", "active_navigation", "stopped_or_parked", "unclear_or_non_process"],
        boundaries=["waiting_vs_moving", "moving_vs_stopped", "traffic_vs_salience"],
    ),
    "affordance": FamilyConfig(
        key="affordance",
        title="Object-Use Affordance",
        scenario_family="object_use_affordance_process",
        question="What stage of the object-use process is shown in this scene?",
        labels=["ready_to_use_setup", "active_use", "used_or_finished", "unclear_or_non_process"],
        boundaries=["setup_vs_active_use", "active_vs_finished_use", "affordance_vs_salience"],
    ),
}


def has_any(text: str, terms: list[str]) -> bool:
    return any(term in text for term in terms)


def contains_any(text: str, terms: list[str]) -> list[str]:
    return [term for term in terms if term in text]


def load_coco_metadata() -> tuple[dict[int, list[str]], dict[int, list[str]], dict[int, str]]:
    with zipfile.ZipFile(ZIP_PATH) as zf:
        caps = json.load(io.TextIOWrapper(zf.open("annotations/captions_val2017.json"), encoding="utf-8"))
        inst = json.load(io.TextIOWrapper(zf.open("annotations/instances_val2017.json"), encoding="utf-8"))
    cat_name = {row["id"]: row["name"] for row in inst["categories"]}
    caps_by_img: dict[int, list[str]] = defaultdict(list)
    cats_by_img: dict[int, list[str]] = defaultdict(list)
    file_by_img: dict[int, str] = {}
    for image in inst["images"]:
        file_by_img[image["id"]] = image["file_name"]
    for ann in caps["annotations"]:
        caps_by_img[ann["image_id"]].append(ann["caption"])
    for ann in inst["annotations"]:
        cats_by_img[ann["image_id"]].append(cat_name[ann["category_id"]])
    return caps_by_img, cats_by_img, file_by_img


def infer_assembly(text: str, cats: set[str]) -> tuple[str, str, str, str, str, str, str, list[str], int] | None:
    terms = ["building", "constructing", "assembling", "putting together", "making", "repairing", "fixing", "installing", "setting up", "tent", "kite", "bench", "chair"]
    triggers = contains_any(text, terms)
    if {"bench", "chair", "kite", "umbrella"} & cats:
        triggers.extend(sorted({"bench", "chair", "kite", "umbrella"} & cats))
    if not triggers:
        return None
    score = len(set(triggers)) + (3 if "person" in cats else 0)
    if has_any(text, ["constructing", "assembling", "putting together", "repairing", "fixing", "installing"]):
        return ("active_assembly", "setup_vs_active", "person/hand", "object/structure", "tool/object contact", "hands/tools are changing the object state", "caption indicates active assembly/repair", triggers, score + 8)
    if has_any(text, ["setting up", "preparing", "about to"]):
        return ("pre_assembly_setup", "setup_vs_active", "person/hand", "object/structure", "near object/materials", "materials or object are present before clear assembly contact", "caption indicates setup before assembly", triggers, score + 5)
    if has_any(text, ["built", "finished", "completed", "structure", "building", "bench", "chair", "tent"]):
        return ("assembled_or_finished", "assembly_vs_finished", "none_or_unclear", "assembled object/structure", "finished object is salient", "object appears as a completed result rather than an active build", "caption/object category suggests finished artifact", triggers, score + 3)
    return ("unclear_or_non_process", "assembly_vs_salience", "none_or_unclear", "salient object", "assembly-related object only", "object salience alone does not ground an assembly stage", "weak assembly salience", triggers, score)


def infer_physical(text: str, cats: set[str]) -> tuple[str, str, str, str, str, str, str, list[str], int] | None:
    terms = ["falling", "fallen", "broken", "cracked", "spilling", "spill", "splashing", "splash", "pouring", "melting", "burning", "fire", "smoke", "wet", "dry", "cut in half", "sliced"]
    triggers = contains_any(text, terms)
    if {"fire hydrant", "cake", "banana", "orange", "wine glass", "cup"} & cats:
        triggers.extend(sorted({"fire hydrant", "cake", "banana", "orange", "wine glass", "cup"} & cats))
    if not triggers:
        return None
    score = len(set(triggers))
    if has_any(text, ["falling", "spilling", "splashing", "pouring", "melting", "burning", "smoke"]):
        return ("active_transition", "stable_vs_active", "gravity/water/fire/tool", "object/material", "visible motion/contact/change cue", "image/caption suggests the transition is currently happening", "active physical change cue", triggers, score + 8)
    if has_any(text, ["fallen", "broken", "cracked", "wet", "cut in half", "sliced"]):
        return ("post_transition_result", "active_vs_result", "none_or_unclear", "changed object/material", "result state visible", "changed state is visible without current transition dynamics", "post-transition result cue", triggers, score + 6)
    if has_any(text, ["dry", "whole", "intact"]):
        return ("pre_transition_stable", "stable_vs_active", "none_or_unclear", "stable object/material", "stable state visible", "object appears before a possible transition", "stable-state cue", triggers, score + 3)
    return ("unclear_or_non_process", "transition_vs_salience", "none_or_unclear", "salient object/material", "physical object is salient", "salience alone does not imply a state transition", "weak physical-transition salience", triggers, score)


def infer_social(text: str, cats: set[str]) -> tuple[str, str, str, str, str, str, str, list[str], int] | None:
    terms = ["wedding", "birthday", "party", "celebration", "ceremony", "parade", "crowd", "audience", "dining", "eating", "meal", "restaurant", "game", "playing", "spectators", "table"]
    triggers = contains_any(text, terms)
    if {"dining table", "cake", "sports ball", "person"} & cats:
        triggers.extend(sorted({"dining table", "cake", "sports ball"} & cats))
    if not triggers or ("person" not in cats and not has_any(text, ["wedding", "birthday", "party", "crowd", "audience"])):
        return None
    score = len(set(triggers)) + (3 if "person" in cats else 0)
    if has_any(text, ["playing", "eating", "celebrating", "ceremony", "parade", "watching", "dining"]):
        return ("active_event", "setup_vs_active_event", "people/group", "event props/space", "people engaged in event roles", "people are actively participating in an event script", "active social/event cue", triggers, score + 7)
    if has_any(text, ["set table", "empty table", "waiting", "preparing", "setup", "decorated"]):
        return ("pre_event_setup", "setup_vs_active_event", "none_or_people_waiting", "event setting/props", "props arranged before participation", "scene shows event affordances before clear participation", "event setup cue", triggers, score + 5)
    if has_any(text, ["leftover", "empty plates", "after", "finished meal", "messy table"]):
        return ("post_event_result", "active_vs_after_event", "none_or_unclear", "event remains", "after-event residue", "objects indicate an event happened but is no longer active", "post-event result cue", triggers, score + 6)
    return ("unclear_or_non_process", "event_vs_salience", "none_or_unclear", "event-like setting", "social props are salient", "scene has event-related objects but no clear script stage", "weak event salience", triggers, score)


def infer_traffic(text: str, cats: set[str]) -> tuple[str, str, str, str, str, str, str, list[str], int] | None:
    terms = ["crossing", "walking across", "traffic", "traffic light", "street", "road", "sidewalk", "intersection", "waiting", "bus stop", "parked", "driving", "riding", "train station"]
    triggers = contains_any(text, terms)
    if {"traffic light", "stop sign", "car", "bus", "truck", "train", "bicycle", "motorcycle"} & cats:
        triggers.extend(sorted({"traffic light", "stop sign", "car", "bus", "truck", "train", "bicycle", "motorcycle"} & cats))
    if not triggers:
        return None
    score = len(set(triggers)) + (2 if "person" in cats else 0)
    if has_any(text, ["crossing", "walking across", "driving", "riding", "moving", "on the road"]):
        return ("active_navigation", "waiting_vs_moving", "person/vehicle", "road/crosswalk/path", "actor/vehicle is in the traffic path", "scene suggests active movement through traffic space", "active navigation cue", triggers, score + 7)
    if has_any(text, ["waiting", "bus stop", "standing at", "train station", "traffic light"]):
        return ("pre_navigation_waiting", "waiting_vs_moving", "person/vehicle", "traffic control/stop area", "waiting or controlled-stop context", "scene suggests waiting before movement", "waiting/setup traffic cue", triggers, score + 5)
    if has_any(text, ["parked", "stopped", "stop sign"]):
        return ("stopped_or_parked", "moving_vs_stopped", "none_or_vehicle", "vehicle/traffic control", "stopped/parked state visible", "vehicle or traffic participant is not actively navigating", "stopped result cue", triggers, score + 6)
    return ("unclear_or_non_process", "traffic_vs_salience", "none_or_unclear", "traffic object/space", "traffic object is salient", "traffic-related object alone does not ground navigation stage", "weak traffic salience", triggers, score)


def infer_affordance(text: str, cats: set[str]) -> tuple[str, str, str, str, str, str, str, list[str], int] | None:
    affordance_cats = {"umbrella", "cell phone", "laptop", "toothbrush", "scissors", "knife", "fork", "spoon", "remote", "keyboard", "mouse", "book", "bottle", "cup"}
    terms = ["using", "holding", "open umbrella", "closed umbrella", "talking on", "typing", "reading", "brushing", "cutting with", "drinking", "eating with", "remote", "keyboard", "phone", "laptop"]
    triggers = contains_any(text, terms)
    if affordance_cats & cats:
        triggers.extend(sorted(affordance_cats & cats))
    if not triggers:
        return None
    score = len(set(triggers)) + (2 if "person" in cats else 0)
    if has_any(text, ["using", "talking on", "typing", "reading", "brushing", "cutting with", "drinking", "eating with", "playing with"]):
        return ("active_use", "setup_vs_active_use", "person/hand", "tool/object", "object is being used with its affordance", "actor-object relation shows active use", "active affordance cue", triggers, score + 8)
    if has_any(text, ["holding", "carrying", "open umbrella", "closed umbrella", "sitting next to", "on a table"]):
        return ("ready_to_use_setup", "setup_vs_active_use", "person/nearby", "tool/object", "object is available or held but use may not be active", "object is ready for use without clear affordance execution", "ready/setup affordance cue", triggers, score + 5)
    if has_any(text, ["empty", "left", "discarded", "used", "finished"]):
        return ("used_or_finished", "active_vs_finished_use", "none_or_unclear", "used object", "post-use state", "object appears after use rather than currently used", "post-use cue", triggers, score + 5)
    return ("unclear_or_non_process", "affordance_vs_salience", "none_or_unclear", "salient affordance object", "object is visible", "affordance-capable object is salient but no use stage is grounded", "weak affordance salience", triggers, score)


INFER = {
    "assembly": infer_assembly,
    "physical": infer_physical,
    "social": infer_social,
    "traffic": infer_traffic,
    "affordance": infer_affordance,
}


def infer_subfamily(family_key: str, text: str, cats: set[str]) -> str:
    if family_key == "assembly":
        if has_any(text, ["repairing", "fixing"]):
            return "repair"
        if has_any(text, ["tent", "kite", "umbrella"]):
            return "setup_object"
        return "constructed_object"
    if family_key == "physical":
        if has_any(text, ["spill", "splash", "pouring"]):
            return "fluid_transition"
        if has_any(text, ["broken", "cracked", "fallen"]):
            return "damage_or_fall"
        return "state_change"
    if family_key == "social":
        if has_any(text, ["wedding", "birthday", "party", "celebration"]):
            return "ceremony_party"
        if has_any(text, ["eating", "meal", "dining"]):
            return "meal_script"
        return "public_event"
    if family_key == "traffic":
        if {"bus", "train"} & cats or has_any(text, ["bus stop", "train station"]):
            return "public_transit"
        if has_any(text, ["crossing", "sidewalk", "intersection"]):
            return "pedestrian_crossing"
        return "road_vehicle"
    if family_key == "affordance":
        for cat in ["umbrella", "cell phone", "laptop", "toothbrush", "scissors", "book", "bottle", "cup"]:
            if cat in cats or cat in text:
                return cat.replace(" ", "_")
        return "general_affordance"
    return "other"


def download_image(image_dir: Path, file_name: str, allow_download: bool) -> Path | None:
    image_dir.mkdir(parents=True, exist_ok=True)
    out_path = image_dir / file_name
    if out_path.exists() and out_path.stat().st_size > 0:
        return out_path
    if not allow_download:
        return None
    url = f"http://images.cocodataset.org/val2017/{file_name}"
    for attempt in range(2):
        try:
            with urllib.request.urlopen(url, timeout=25, context=SSL_CONTEXT) as response:
                out_path.write_bytes(response.read())
            if out_path.exists() and out_path.stat().st_size > 0:
                return out_path
        except Exception:
            time.sleep(2**attempt)
    return None


def build_family(family: FamilyConfig, limit: int, allow_download: bool) -> list[Candidate]:
    caps_by_img, cats_by_img, file_by_img = load_coco_metadata()
    infer = INFER[family.key]
    image_dir = DATA_DIR / f"task3_{family.key}_coco_images"
    metas = []
    for image_id, captions in caps_by_img.items():
        text = " ".join(captions[:5]).lower()
        cats = set(cats_by_img.get(image_id, []))
        inferred = infer(text, cats)
        if inferred:
            metas.append((image_id, inferred))
    metas.sort(key=lambda row: (-row[1][-1], row[0]))

    selected: list[Candidate] = []
    per_label: defaultdict[str, int] = defaultdict(int)
    target_per_label = max(8, limit // max(1, len(family.labels)))
    for image_id, inferred in metas:
        label = inferred[0]
        if per_label[label] >= target_per_label and len(selected) < limit - 8:
            continue
        file_name = file_by_img.get(image_id, f"{image_id:012d}.jpg")
        path = download_image(image_dir, file_name, allow_download)
        if not path:
            continue
        text = " ".join(caps_by_img.get(image_id, [])[:5]).lower()
        cats = set(cats_by_img.get(image_id, []))
        selected.append(
            Candidate(
                family=family,
                image_id=image_id,
                image_path=str(path.resolve()),
                captions=caps_by_img.get(image_id, [])[:5],
                categories=sorted(cats),
                suggested_label=label,
                suggested_boundary_family=inferred[1],
                suggested_actor=inferred[2],
                suggested_object=inferred[3],
                suggested_relation=inferred[4],
                suggested_state_cues=inferred[5],
                suggestion_reason=inferred[6],
                trigger_words=inferred[7],
                priority_score=inferred[8],
                subfamily=infer_subfamily(family.key, text, cats),
            )
        )
        per_label[label] += 1
        if len(selected) >= limit:
            break
    return selected


def default_neighbor(family: FamilyConfig, label: str) -> str:
    labels = family.labels
    if label == labels[0]:
        return labels[1]
    if label == labels[1]:
        return labels[0]
    if label == labels[2]:
        return labels[1]
    return labels[1]


def write_csv(path: Path, rows: list[Candidate]) -> None:
    headers = [
        "id",
        "image_path",
        "keep",
        "question",
        "gold_step_label",
        "neighbor_step_label",
        "boundary_family",
        "why_not_neighbor",
        "visible_actor",
        "visible_object",
        "visible_relation",
        "visible_state_cues",
        "issue_type",
        "notes",
        "source_dataset",
        "source_image_id",
        "suggested_label",
        "suggested_boundary_family",
        "suggested_actor",
        "suggested_object",
        "suggested_relation",
        "suggested_state_cues",
        "suggestion_reason",
        "captions",
        "categories",
        "trigger_words",
        "priority_score",
        "subfamily",
        "scenario_family",
    ]
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=headers)
        writer.writeheader()
        for idx, row in enumerate(rows, 1):
            writer.writerow(
                {
                    "id": f"task3_{row.family.key}_{idx:04d}_{row.image_id}",
                    "image_path": row.image_path,
                    "keep": "",
                    "question": row.family.question,
                    "gold_step_label": "",
                    "neighbor_step_label": "",
                    "boundary_family": "",
                    "why_not_neighbor": "",
                    "visible_actor": "",
                    "visible_object": "",
                    "visible_relation": "",
                    "visible_state_cues": "",
                    "issue_type": "",
                    "notes": "",
                    "source_dataset": "COCO val2017",
                    "source_image_id": row.image_id,
                    "suggested_label": row.suggested_label,
                    "suggested_boundary_family": row.suggested_boundary_family,
                    "suggested_actor": row.suggested_actor,
                    "suggested_object": row.suggested_object,
                    "suggested_relation": row.suggested_relation,
                    "suggested_state_cues": row.suggested_state_cues,
                    "suggestion_reason": row.suggestion_reason,
                    "captions": " || ".join(row.captions),
                    "categories": ", ".join(row.categories),
                    "trigger_words": ", ".join(sorted(set(row.trigger_words))),
                    "priority_score": row.priority_score,
                    "subfamily": row.subfamily,
                    "scenario_family": row.family.scenario_family,
                }
            )


def write_html(path: Path, family: FamilyConfig, rows: list[Candidate]) -> None:
    cards = []
    for idx, row in enumerate(rows, 1):
        image_src = f"file://{row.image_path}"
        cards.append(
            f"""
            <article class="card">
              <h2>{idx:03d} / COCO {row.image_id}</h2>
              <img src="{html.escape(image_src)}" />
              <p><b>Suggested:</b> {html.escape(row.suggested_label)} / {html.escape(row.suggested_boundary_family)}</p>
              <p><b>Subfamily:</b> {html.escape(row.subfamily)}</p>
              <p><b>Actor-object:</b> {html.escape(row.suggested_actor)} -> {html.escape(row.suggested_object)}</p>
              <p><b>Relation:</b> {html.escape(row.suggested_relation)}</p>
              <p><b>State cues:</b> {html.escape(row.suggested_state_cues)}</p>
              <p><b>Reason:</b> {html.escape(row.suggestion_reason)}</p>
              <p><b>Captions:</b> {html.escape(" || ".join(row.captions))}</p>
              <p><b>Categories:</b> {html.escape(", ".join(row.categories))}</p>
            </article>
            """
        )
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(
        f"""<!doctype html>
<html>
<head>
  <meta charset="utf-8" />
  <title>Task3 {html.escape(family.title)} Review</title>
  <style>
    body {{ font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", sans-serif; margin: 24px; background: #f7f7f5; color: #1d1d1f; }}
    .grid {{ display: grid; grid-template-columns: repeat(auto-fill, minmax(360px, 1fr)); gap: 18px; }}
    .card {{ background: white; border: 1px solid #d9d9d6; border-radius: 8px; padding: 14px; }}
    img {{ width: 100%; max-height: 320px; object-fit: contain; background: #eee; }}
    h1 {{ margin-top: 0; }}
    h2 {{ font-size: 16px; margin: 0 0 10px; }}
    p {{ font-size: 13px; line-height: 1.35; }}
  </style>
</head>
<body>
  <h1>Task3 {html.escape(family.title)} Review</h1>
  <p>Use the xlsx file for annotation. This page is for visual inspection.</p>
  <p><b>Answer space:</b> {html.escape(", ".join(family.labels))}</p>
  <p><b>Boundary families:</b> {html.escape(", ".join(family.boundaries))}</p>
  <div class="grid">{''.join(cards)}</div>
</body>
</html>
""",
        encoding="utf-8",
    )


def write_xlsx(csv_path: Path, xlsx_path: Path, family: FamilyConfig) -> None:
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    rows = list(csv.DictReader(csv_path.open("r", encoding="utf-8")))
    wb = Workbook()
    ws = wb.active
    ws.title = f"{family.key}_review"
    headers = list(rows[0].keys()) if rows else []
    ws.append(headers)
    for row in rows:
        ws.append([row.get(h, "") for h in headers])

    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="DEEAF6")
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    widths = {"A": 28, "B": 58, "C": 10, "D": 48, "E": 26, "F": 26, "G": 30, "H": 46, "I": 20, "J": 24, "K": 34, "L": 40, "M": 24, "N": 32, "W": 70, "X": 70}
    for col, width in widths.items():
        ws.column_dimensions[col].width = width
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    header_to_col = {cell.value: cell.column_letter for cell in ws[1]}
    options = {
        "keep": ["yes", "no"],
        "gold_step_label": family.labels,
        "neighbor_step_label": family.labels + ["none"],
        "boundary_family": family.boundaries,
        "issue_type": ["ok", "bad_image", "not_process", "label_unclear", "duplicate", "too_hard", "other"],
    }
    for header, values in options.items():
        col = header_to_col.get(header)
        if not col:
            continue
        dv = DataValidation(type="list", formula1=f'"{",".join(values)}"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"{col}2:{col}{ws.max_row}")

    comments = {
        "keep": "yes means this image is visually suitable for the process benchmark; no excludes it.",
        "gold_step_label": "Correct visible process stage. Use suggested_label if it is right.",
        "neighbor_step_label": "Tempting but wrong neighboring stage.",
        "boundary_family": "The process boundary being tested.",
        "why_not_neighbor": "One short visual reason why the neighbor is wrong.",
        "visible_state_cues": "Observable visual cues only.",
    }
    for header, text in comments.items():
        col = header_to_col.get(header)
        if col:
            ws[f"{col}1"].comment = Comment(text, "Codex")
    xlsx_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(xlsx_path)


def write_bootstrap_jsonl(path: Path, rows: list[Candidate]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8") as f:
        for idx, row in enumerate(rows, 1):
            item = {
                "id": f"task3_{row.family.key}_{idx:04d}_{row.image_id}",
                "image_id": str(row.image_id),
                "source": f"task3_{row.family.key}_bootstrap",
                "source_dataset": "COCO val2017",
                "task": f"task3_{row.family.scenario_family}",
                "scenario_family": row.family.scenario_family,
                "subfamily": row.subfamily,
                "underspecified_question": row.family.question,
                "original_question": row.family.question,
                "answer_space": row.family.labels,
                "gold_label": row.suggested_label,
                "expected_answer": row.suggested_label,
                "answer": row.suggested_label,
                "label_source": "suggested_bootstrap",
                "boundary_family": row.suggested_boundary_family,
                "boundary_family_hint": row.suggested_boundary_family,
                "boundary_source": "suggested_bootstrap",
                "neighbor_label": default_neighbor(row.family, row.suggested_label),
                "why_not_neighbor": row.suggestion_reason,
                "visible_actor": row.suggested_actor,
                "visible_object": row.suggested_object,
                "visible_relation": row.suggested_relation,
                "visible_state_cues": row.suggested_state_cues,
                "suggested_label": row.suggested_label,
                "suggested_boundary_family": row.suggested_boundary_family,
                "suggestion_reason": row.suggestion_reason,
                "captions": " || ".join(row.captions),
                "categories": " | ".join(row.categories),
                "image_path": row.image_path,
            }
            f.write(json.dumps(item, ensure_ascii=False) + "\n")


def summarize(rows: list[Candidate]) -> dict[str, object]:
    out: dict[str, object] = {"n": len(rows), "label_distribution": {}, "boundary_distribution": {}, "subfamily_distribution": {}}
    for row in rows:
        for key, value in [
            ("label_distribution", row.suggested_label),
            ("boundary_distribution", row.suggested_boundary_family),
            ("subfamily_distribution", row.subfamily),
        ]:
            dist = out[key]
            assert isinstance(dist, dict)
            dist[value] = dist.get(value, 0) + 1
    return out


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--families", nargs="+", default=list(FAMILIES))
    parser.add_argument("--limit", type=int, default=60)
    parser.add_argument("--no-download", action="store_true")
    args = parser.parse_args()

    all_summary = {}
    for key in args.families:
        family = FAMILIES[key]
        rows = build_family(family, args.limit, allow_download=not args.no_download)
        base = f"task3_{key}_candidates_v1_review"
        csv_path = ROOT / "review" / f"{base}.csv"
        xlsx_path = ROOT / "review" / f"{base}.xlsx"
        html_path = ROOT / "review" / f"{base}.html"
        jsonl_path = ROOT / "data" / f"task3_{key}_eval_v1_bootstrap.jsonl"
        write_csv(csv_path, rows)
        write_xlsx(csv_path, xlsx_path, family)
        write_html(html_path, family, rows)
        write_bootstrap_jsonl(jsonl_path, rows)
        all_summary[key] = {
            **summarize(rows),
            "csv": str(csv_path),
            "xlsx": str(xlsx_path),
            "html": str(html_path),
            "bootstrap_jsonl": str(jsonl_path),
        }
    summary_path = ROOT / "runs" / "task3_remaining_process_candidates_v1_summary_2026-07-02.json"
    summary_path.parent.mkdir(parents=True, exist_ok=True)
    summary_path.write_text(json.dumps(all_summary, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"summary": str(summary_path), "families": all_summary}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
