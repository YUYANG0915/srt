#!/usr/bin/env python3
"""Build Task3 P5 cleaning evaluation items from the reviewed XLSX/CSV."""

from __future__ import annotations

import argparse
import csv
import json
from pathlib import Path
from typing import Any


LABELS = {"pre_cleaning_setup", "active_cleaning", "cleaned_or_finished", "unclear_or_non_process"}


def truthy(value: Any) -> bool:
    return str(value or "").strip().lower() in {"yes", "y", "true", "1", "keep"}


def clean(value: Any) -> str:
    return str(value or "").strip()


def read_csv(path: Path) -> list[dict[str, Any]]:
    with path.open("r", encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


def read_xlsx(path: Path) -> list[dict[str, Any]]:
    from openpyxl import load_workbook

    workbook = load_workbook(path, data_only=True)
    sheet = workbook["Review"]
    headers = [cell.value for cell in sheet[1]]
    rows: list[dict[str, Any]] = []
    for row_idx in range(2, sheet.max_row + 1):
        row = {
            str(header): sheet.cell(row_idx, col_idx).value
            for col_idx, header in enumerate(headers, start=1)
            if header
        }
        if clean(row.get("id")):
            rows.append(row)
    return rows


def read_rows(path: Path) -> list[dict[str, Any]]:
    if path.suffix.lower() == ".xlsx":
        return read_xlsx(path)
    return read_csv(path)


def choose_label(row: dict[str, Any]) -> tuple[str, str]:
    gold = clean(row.get("gold_label"))
    if gold in LABELS:
        return gold, "human_gold"
    suggested = clean(row.get("suggested_label"))
    if suggested in LABELS:
        return suggested, "suggested_fallback_after_keep_only"
    return "unclear_or_non_process", "fallback_unclear"


def choose_boundary(row: dict[str, Any]) -> tuple[str, str]:
    boundary = clean(row.get("boundary_family"))
    if boundary:
        return boundary, "human_boundary"
    suggested = clean(row.get("suggested_boundary_family"))
    if suggested:
        return suggested, "suggested_fallback_after_keep_only"
    return "other", "fallback_other"


def default_neighbor(label: str) -> str:
    if label == "pre_cleaning_setup":
        return "active_cleaning"
    if label == "active_cleaning":
        return "pre_cleaning_setup"
    if label == "cleaned_or_finished":
        return "active_cleaning"
    return "active_cleaning"


def build_item(row: dict[str, Any]) -> dict[str, Any]:
    label, label_source = choose_label(row)
    boundary, boundary_source = choose_boundary(row)
    question = clean(row.get("question")) or "What cleaning stage is shown in this scene?"
    image_path = str(Path(clean(row["image_path"])).expanduser().resolve())
    tool = clean(row.get("visible_tool")) or clean(row.get("suggested_tool"))
    target = clean(row.get("visible_target")) or clean(row.get("suggested_target"))
    state_cues = clean(row.get("visible_state_cues")) or clean(row.get("suggested_state_cues")) or clean(row.get("suggestion_reason"))
    neighbor = clean(row.get("neighbor_label"))
    if not neighbor or neighbor not in LABELS:
        neighbor = default_neighbor(label)

    return {
        "id": clean(row["id"]),
        "image_id": clean(row.get("source_image_id")) or Path(image_path).stem,
        "source": "task3_cleaning_reviewed",
        "source_dataset": clean(row.get("source_dataset")) or "coco_val2017_downloaded",
        "task": "task3_cleaning_household_state_change",
        "scenario_family": "cleaning_household_state_change",
        "underspecified_question": question,
        "original_question": question,
        "answer_space": [
            "pre_cleaning_setup",
            "active_cleaning",
            "cleaned_or_finished",
            "unclear_or_non_process",
        ],
        "gold_label": label,
        "expected_answer": label,
        "answer": label,
        "label_source": label_source,
        "boundary_family": boundary,
        "boundary_family_hint": boundary,
        "boundary_source": boundary_source,
        "neighbor_label": neighbor,
        "why_not_neighbor": clean(row.get("why_not_neighbor")),
        "visible_actor": clean(row.get("visible_actor")),
        "visible_tool": tool,
        "visible_target": target,
        "visible_state_cues": state_cues,
        "suggested_label": clean(row.get("suggested_label")),
        "suggested_boundary_family": clean(row.get("suggested_boundary_family")),
        "suggestion_reason": clean(row.get("suggestion_reason")),
        "captions": clean(row.get("captions")),
        "categories": clean(row.get("categories")),
        "issue_type": clean(row.get("issue_type")),
        "notes": clean(row.get("notes")),
        "image_path": image_path,
    }


def write_jsonl(path: Path, rows: list[dict[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("".join(json.dumps(row, ensure_ascii=False) + "\n" for row in rows), encoding="utf-8")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--review", default="review/task3_cleaning_candidates_v1_review.xlsx")
    parser.add_argument("--output", default="data/task3_cleaning_eval_v1_reviewed.jsonl")
    parser.add_argument("--summary", default="runs/task3_cleaning_eval_v1_reviewed_summary_2026-06-30.json")
    args = parser.parse_args()

    rows = [row for row in read_rows(Path(args.review)) if truthy(row.get("keep"))]
    if not rows:
        raise ValueError("No keep=yes rows found.")
    items = [build_item(row) for row in rows]
    write_jsonl(Path(args.output), items)

    summary: dict[str, Any] = {
        "review": str(Path(args.review).resolve()),
        "output": str(Path(args.output).resolve()),
        "n": len(items),
        "label_distribution": {},
        "boundary_distribution": {},
        "label_source_distribution": {},
    }
    for item in items:
        summary["label_distribution"][item["gold_label"]] = summary["label_distribution"].get(item["gold_label"], 0) + 1
        summary["boundary_distribution"][item["boundary_family"]] = summary["boundary_distribution"].get(item["boundary_family"], 0) + 1
        summary["label_source_distribution"][item["label_source"]] = summary["label_source_distribution"].get(item["label_source"], 0) + 1
    Path(args.summary).parent.mkdir(parents=True, exist_ok=True)
    Path(args.summary).write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
