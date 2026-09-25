#!/usr/bin/env python3
"""Apply focused correction XLSX to Task3 P6 craft items."""

from __future__ import annotations

import argparse
import json
from pathlib import Path
from typing import Any


LABELS = {
    "pre_creation_setup",
    "active_creation",
    "active_assembly",
    "finished_artifact",
    "unclear_or_non_creation",
}
BOUNDARIES = {
    "setup_vs_active",
    "active_vs_finished",
    "assembly_vs_finished",
    "creation_vs_salience",
}


def clean(value: Any) -> str:
    return str(value or "").strip()


def truthy(value: Any) -> bool:
    return clean(value).lower() in {"yes", "y", "true", "1", "keep"}


def load_jsonl(path: Path) -> list[dict[str, Any]]:
    return [json.loads(line) for line in path.open("r", encoding="utf-8") if line.strip()]


def read_xlsx(path: Path) -> list[dict[str, Any]]:
    from openpyxl import load_workbook

    wb = load_workbook(path, data_only=True)
    ws = wb.active
    headers = [cell.value for cell in ws[1]]
    rows = []
    for row_idx in range(2, ws.max_row + 1):
        row = {
            str(header): ws.cell(row_idx, col_idx).value
            for col_idx, header in enumerate(headers, start=1)
            if header
        }
        if clean(row.get("id")):
            rows.append(row)
    return rows


def default_neighbor(label: str) -> str:
    if label == "pre_creation_setup":
        return "active_creation"
    if label == "active_creation":
        return "pre_creation_setup"
    if label == "active_assembly":
        return "finished_artifact"
    if label == "finished_artifact":
        return "active_creation"
    return "active_creation"


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--base-items", default="data/task3_craft_eval_v1_reviewed.jsonl")
    parser.add_argument("--correction", default="review/task3_craft_focused_correction_v1.xlsx")
    parser.add_argument("--output", default="data/task3_craft_eval_v1_corrected.jsonl")
    parser.add_argument("--summary", default="runs/task3_craft_eval_v1_corrected_summary_2026-07-01.json")
    args = parser.parse_args()

    items = {row["id"]: row for row in load_jsonl(Path(args.base_items))}
    corrections = {clean(row["id"]): row for row in read_xlsx(Path(args.correction))}

    corrected_items: list[dict[str, Any]] = []
    dropped = []
    for item_id, item in items.items():
        row = corrections.get(item_id, {})
        keep_value = clean(row.get("keep"))
        if keep_value and not truthy(keep_value):
            dropped.append(item_id)
            continue
        new_item = dict(item)
        corrected_label = clean(row.get("corrected_gold_step_label"))
        corrected_boundary = clean(row.get("corrected_boundary_family"))
        corrected_neighbor = clean(row.get("corrected_neighbor_step_label"))

        if corrected_label in LABELS:
            new_item["gold_label"] = corrected_label
            new_item["expected_answer"] = corrected_label
            new_item["answer"] = corrected_label
            new_item["label_source"] = "focused_human_correction"
        else:
            new_item["label_source"] = item.get("label_source", "suggested_fallback_after_keep_only")

        if corrected_boundary in BOUNDARIES:
            new_item["boundary_family"] = corrected_boundary
            new_item["boundary_family_hint"] = corrected_boundary
            new_item["boundary_source"] = "focused_human_correction"
        else:
            new_item["boundary_source"] = item.get("boundary_source", "suggested_fallback_after_keep_only")

        if corrected_neighbor in LABELS:
            new_item["neighbor_label"] = corrected_neighbor
        elif clean(new_item.get("neighbor_label")) not in LABELS:
            new_item["neighbor_label"] = default_neighbor(str(new_item["gold_label"]))

        for key in ("why_not_neighbor", "issue_type", "notes"):
            value = clean(row.get(key))
            if value:
                new_item[key] = value
        corrected_items.append(new_item)

    Path(args.output).parent.mkdir(parents=True, exist_ok=True)
    Path(args.output).write_text("".join(json.dumps(row, ensure_ascii=False) + "\n" for row in corrected_items), encoding="utf-8")

    summary: dict[str, Any] = {
        "base_items": str(Path(args.base_items).resolve()),
        "correction": str(Path(args.correction).resolve()),
        "output": str(Path(args.output).resolve()),
        "n": len(corrected_items),
        "dropped": len(dropped),
        "label_distribution": {},
        "boundary_distribution": {},
        "label_source_distribution": {},
    }
    for item in corrected_items:
        summary["label_distribution"][item["gold_label"]] = summary["label_distribution"].get(item["gold_label"], 0) + 1
        summary["boundary_distribution"][item["boundary_family"]] = summary["boundary_distribution"].get(item["boundary_family"], 0) + 1
        summary["label_source_distribution"][item["label_source"]] = summary["label_source_distribution"].get(item["label_source"], 0) + 1

    Path(args.summary).parent.mkdir(parents=True, exist_ok=True)
    Path(args.summary).write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
