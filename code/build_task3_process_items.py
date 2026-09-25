#!/usr/bin/env python3
"""Convert a reviewed SRT process-family CSV/XLSX into eval JSONL."""

from __future__ import annotations

import argparse
import csv
import json
from pathlib import Path
from typing import Any

from build_task3_remaining_process_candidates import FAMILIES, default_neighbor


def truthy(value: Any) -> bool:
    return str(value or "").strip().lower() in {"yes", "y", "true", "1", "keep"}


def clean(value: Any) -> str:
    return str(value or "").strip()


def read_csv(path: Path) -> list[dict[str, Any]]:
    with path.open("r", encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


def read_xlsx(path: Path) -> list[dict[str, Any]]:
    from openpyxl import load_workbook

    workbook = load_workbook(path, data_only=True)
    sheet = workbook.active
    headers = [cell.value for cell in sheet[1]]
    rows: list[dict[str, Any]] = []
    for row_idx in range(2, sheet.max_row + 1):
        row = {
            str(header): sheet.cell(row_idx, col_idx).value
            for col_idx, header in enumerate(headers, start=1)
            if header
        }
        if clean(row.get("id")):
            rows.append(row)
    return rows


def read_rows(path: Path) -> list[dict[str, Any]]:
    if path.suffix.lower() == ".xlsx":
        return read_xlsx(path)
    return read_csv(path)


def choose_label(row: dict[str, Any], labels: list[str]) -> tuple[str, str]:
    gold = clean(row.get("gold_step_label"))
    if gold in labels:
        return gold, "human_gold"
    suggested = clean(row.get("suggested_label"))
    if suggested in labels:
        return suggested, "suggested_fallback_after_keep_only"
    return labels[-1], "fallback_unclear"


def choose_boundary(row: dict[str, Any], boundaries: list[str]) -> tuple[str, str]:
    boundary = clean(row.get("boundary_family"))
    if boundary in boundaries:
        return boundary, "human_boundary"
    suggested = clean(row.get("suggested_boundary_family"))
    if suggested in boundaries:
        return suggested, "suggested_fallback_after_keep_only"
    return boundaries[-1], "fallback_salience"


def build_item(row: dict[str, Any], family_key: str) -> dict[str, Any]:
    family = FAMILIES[family_key]
    label, label_source = choose_label(row, family.labels)
    boundary, boundary_source = choose_boundary(row, family.boundaries)
    image_path = str(Path(clean(row["image_path"])).expanduser().resolve())
    neighbor = clean(row.get("neighbor_step_label"))
    if neighbor not in family.labels:
        neighbor = default_neighbor(family, label)
    question = clean(row.get("question")) or family.question
    return {
        "id": clean(row["id"]),
        "image_id": clean(row.get("source_image_id")) or Path(image_path).stem,
        "source": f"task3_{family_key}_reviewed",
        "source_dataset": clean(row.get("source_dataset")) or "COCO val2017",
        "task": f"task3_{family.scenario_family}",
        "scenario_family": family.scenario_family,
        "subfamily": clean(row.get("subfamily")),
        "underspecified_question": question,
        "original_question": question,
        "answer_space": family.labels,
        "gold_label": label,
        "expected_answer": label,
        "answer": label,
        "label_source": label_source,
        "boundary_family": boundary,
        "boundary_family_hint": boundary,
        "boundary_source": boundary_source,
        "neighbor_label": neighbor,
        "why_not_neighbor": clean(row.get("why_not_neighbor")) or clean(row.get("suggestion_reason")),
        "visible_actor": clean(row.get("visible_actor")) or clean(row.get("suggested_actor")),
        "visible_object": clean(row.get("visible_object")) or clean(row.get("suggested_object")),
        "visible_relation": clean(row.get("visible_relation")) or clean(row.get("suggested_relation")),
        "visible_state_cues": clean(row.get("visible_state_cues")) or clean(row.get("suggested_state_cues")),
        "suggested_label": clean(row.get("suggested_label")),
        "suggested_boundary_family": clean(row.get("suggested_boundary_family")),
        "suggestion_reason": clean(row.get("suggestion_reason")),
        "captions": clean(row.get("captions")),
        "categories": clean(row.get("categories")),
        "issue_type": clean(row.get("issue_type")),
        "notes": clean(row.get("notes")),
        "image_path": image_path,
    }


def write_jsonl(path: Path, rows: list[dict[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("".join(json.dumps(row, ensure_ascii=False) + "\n" for row in rows), encoding="utf-8")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--family", choices=sorted(FAMILIES), required=True)
    parser.add_argument("--review", required=True)
    parser.add_argument("--output", required=True)
    parser.add_argument("--summary", required=True)
    args = parser.parse_args()

    kept = [row for row in read_rows(Path(args.review)) if truthy(row.get("keep"))]
    if not kept:
        raise ValueError("No keep=yes rows found.")
    items = [build_item(row, args.family) for row in kept]
    write_jsonl(Path(args.output), items)

    summary: dict[str, Any] = {
        "family": args.family,
        "review": str(Path(args.review).resolve()),
        "output": str(Path(args.output).resolve()),
        "n": len(items),
        "label_distribution": {},
        "boundary_distribution": {},
        "subfamily_distribution": {},
        "label_source_distribution": {},
        "boundary_source_distribution": {},
    }
    for item in items:
        for key, field in [
            ("label_distribution", "gold_label"),
            ("boundary_distribution", "boundary_family"),
            ("subfamily_distribution", "subfamily"),
            ("label_source_distribution", "label_source"),
            ("boundary_source_distribution", "boundary_source"),
        ]:
            value = str(item.get(field) or "")
            summary[key][value] = summary[key].get(value, 0) + 1
    Path(args.summary).parent.mkdir(parents=True, exist_ok=True)
    Path(args.summary).write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
